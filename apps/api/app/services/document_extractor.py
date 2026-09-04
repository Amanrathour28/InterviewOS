"""
Document Text Extraction Pipeline — Phase 13.

Extracts text and page/section provenance from PDF and DOCX files.
Uses PyMuPDF (fitz) for PDF and python-docx for DOCX, with pure text fallbacks.
"""

import io
import logging
from typing import Any, Dict, List, Optional, Tuple

logger = logging.getLogger("interviewos.document_extractor")


class ExtractedPage:
    def __init__(self, page_number: int, text: str):
        self.page_number = page_number
        self.text = text.strip()


class ExtractionResult:
    def __init__(
        self,
        full_text: str,
        pages: List[ExtractedPage],
        mime_type: str,
        page_count: int,
        metadata: Optional[Dict[str, Any]] = None,
    ):
        self.full_text = full_text
        self.pages = pages
        self.mime_type = mime_type
        self.page_count = page_count
        self.metadata = metadata or {}


def extract_text_from_pdf(file_bytes: bytes) -> ExtractionResult:
    """Extracts text page by page from PDF using PyMuPDF (fitz)."""
    try:
        import fitz  # PyMuPDF

        doc = fitz.open(stream=file_bytes, filetype="pdf")
        pages: List[ExtractedPage] = []
        full_text_chunks: List[str] = []

        for idx, page in enumerate(doc):
            page_text = page.get_text("text") or ""
            pages.append(ExtractedPage(page_number=idx + 1, text=page_text))
            if page_text.strip():
                full_text_chunks.append(f"--- Page {idx + 1} ---\n{page_text.strip()}")

        meta = {
            "format": "PDF",
            "page_count": len(doc),
            "title": doc.metadata.get("title", ""),
            "author": doc.metadata.get("author", ""),
        }
        doc.close()

        full_text = "\n\n".join(full_text_chunks)
        return ExtractionResult(
            full_text=full_text,
            pages=pages,
            mime_type="application/pdf",
            page_count=len(pages),
            metadata=meta,
        )
    except Exception as e:
        logger.warning("PyMuPDF extraction failed or not available: %s. Attempting fallback.", e)
        # Fallback: simple text decode
        decoded = file_bytes.decode("utf-8", errors="ignore")
        return ExtractionResult(
            full_text=decoded,
            pages=[ExtractedPage(1, decoded)],
            mime_type="application/pdf",
            page_count=1,
            metadata={"fallback": True, "error": str(e)},
        )


def extract_text_from_docx(file_bytes: bytes) -> ExtractionResult:
    """Extracts paragraphs and tables from DOCX using python-docx."""
    try:
        import docx

        doc_io = io.BytesIO(file_bytes)
        doc = docx.Document(doc_io)
        paragraphs: List[str] = []

        for p in doc.paragraphs:
            if p.text.strip():
                paragraphs.append(p.text.strip())

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    paragraphs.append(" | ".join(row_cells))

        full_text = "\n\n".join(paragraphs)
        pages = [ExtractedPage(page_number=1, text=full_text)]

        meta = {
            "format": "DOCX",
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
        }

        return ExtractionResult(
            full_text=full_text,
            pages=pages,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            page_count=1,
            metadata=meta,
        )
    except Exception as e:
        logger.warning("python-docx extraction failed: %s. Attempting fallback.", e)
        decoded = file_bytes.decode("utf-8", errors="ignore")
        return ExtractionResult(
            full_text=decoded,
            pages=[ExtractedPage(1, decoded)],
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            page_count=1,
            metadata={"fallback": True, "error": str(e)},
        )


def extract_document_text(file_bytes: bytes, file_name: str, mime_type: str) -> ExtractionResult:
    """Universal document extractor dispatching by mime type and extension."""
    ext = file_name.lower().split(".")[-1] if "." in file_name else ""
    mime = mime_type.lower()

    if "pdf" in mime or ext == "pdf":
        return extract_text_from_pdf(file_bytes)
    elif "docx" in mime or "word" in mime or ext in ("docx", "doc"):
        return extract_text_from_docx(file_bytes)
    else:
        # Default plain text extraction
        decoded = file_bytes.decode("utf-8", errors="replace")
        return ExtractionResult(
            full_text=decoded,
            pages=[ExtractedPage(1, decoded)],
            mime_type=mime_type or "text/plain",
            page_count=1,
            metadata={"format": "PLAIN_TEXT"},
        )
